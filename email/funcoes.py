"""
Módulo de funções para o sistema de gerenciamento e envio de e-mails.
Bibliotecas utilizadas:
- os: Manipulação de caminhos e verificação de arquivos no sistema operacional.
- pathlib: Manipulação moderna de caminhos compatível entre Windows e Linux.
- openpyxl: Criação, leitura e edição de planilhas no formato Excel (.xlsx).
- smtplib: Conexão e envio de e-mails via servidores SMTP.
- email.mime: Estruturação de mensagens MIME (texto, HTML, anexos, multipart).
- datetime: Manipulação de datas e horários para registro de logs e histórico.
"""

import os
import smtplib
import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders

# Variáveis constantes globais para o módulo
ARQUIVO_EXCEL = "destinatarios.xlsx"
HISTORICO_ENVIO = "historico.txt"
CONFIG_SMTP = "config.txt"

# Constantes de elementos decorativos usados para UI / UX
OK = "[✔]"
ERRO = "[✕]"
ATENCAO = "[⚠]"
LI = "➤"


def criarPlanilhaDestinatarios(caminhoArquivo):
    """
    Cria uma nova planilha Excel contendo a tabela estruturada de destinatários.
    Parâmetros:
        - caminhoArquivo (Path): Caminho completo onde o arquivo .xlsx será salvo.
    Métodos e Objetos usados:
        - Workbook(): Instancia uma nova pasta de trabalho do openpyxl.
        - active: Seleciona a aba padrão ativa na planilha.
        - append(): Adiciona uma linha de dados de uma só vez na planilha.
        - save(): Salva a pasta de trabalho no disco com o caminho especificado.
    """
    try:
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Destinatarios"
        
        # Cabeçalho da tabela
        planilha.append(["Nome", "E-mail", "Empresa", "Mensagem Personalizada"])
        
        # Registros de exemplo
        planilha.append(["AMFDZN", "amfdzn@gmail.com", "Acelio Filho Marketing e Design", "Olá Acelio, temos novidades no nosso sistema de automações!"])
        planilha.append(["AMF", "amf.cadastros@gmail.com", "Lead AMF", "Olá Lead Amf, confira nossa nova proposta."])
        planilha.append(["Camiseta Loca", "camiseta.loca@gmail.com", "Camiseta loca", "Olá Camiseta Loca, confira nossa nova proposta."])
        planilha.append(["Banda Xevi50", "xevi50@gmail.com", "Xevi50", "Olá Xevi50, confira nossa nova proposta."])
        
        workbook.save(caminhoArquivo)
        print(f"{OK} Planilha de destinatários criada com sucesso em: {caminhoArquivo}")
    except Exception as erro:
        print(f"{ERRO} Erro ao criar a planilha: {erro}")


def verificarArquivoExcel(caminhoArquivo):
    """
    Verifica se o arquivo Excel de destinatários existe na pasta do exercício.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo a ser verificado.
    Retorno:
        - bool: True se o arquivo existir, False caso contrário.
    Métodos e Objetos usados:
        - is_file(): Método do pathlib que verifica se o caminho aponta para um arquivo existente.
    """
    if caminhoArquivo.is_file():
        print(f"{OK} O arquivo '{caminhoArquivo.name}' existe e está pronto para uso.")
        return True
    else:
        print(f"{ATENCAO} O arquivo '{caminhoArquivo.name}' não foi encontrado. Crie a base primeiro (Opção 0).")
        return False


def abrirPlanilhaExcel(caminhoArquivo):
    """
    Abre uma planilha Excel existente e retorna o objeto carregado para uso.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo .xlsx a ser aberto.
    Retorno:
        - Workbook / None: Retorna o objeto da planilha carregada ou None se houver erro.
    Métodos e Objetos usados:
        - load_workbook(): Carrega uma planilha Excel existente a partir do caminho fornecido.
    """
    if not verificarArquivoExcel(caminhoArquivo):
        return None
    
    try:
        workbook = load_workbook(caminhoArquivo)
        print(f"{OK} Planilha '{caminhoArquivo.name}' aberta com sucesso.")
        return workbook
    except Exception as erro:
        print(f"{ERRO} Erro ao abrir a planilha: {erro}")
        return None


def listarAbasPlanilha(caminhoArquivo):
    """
    Apresenta todas as abas (sheets) existentes dentro do arquivo Excel com otimização de UX.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo .xlsx.
    Métodos e Objetos usados:
        - sheetnames: Atributo do openpyxl que retorna uma lista com o nome de todas as abas da planilha.
        - enumerate(): Itera sobre a lista de abas gerando um índice numérico correspondente.
    """
    workbook = abrirPlanilhaExcel(caminhoArquivo)
    if workbook:
        print(f"\nABAS DA PLANILHA [{ARQUIVO_EXCEL}]:")
        for indice, aba in enumerate(workbook.sheetnames):
            print(f"{LI} Aba Nª{indice}: {aba}")


def visualizarDestinatarios(caminhoArquivo):
    """
    Visualiza e exibe no console todos os destinatários cadastrados na planilha.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo .xlsx.
    Métodos e Objetos usados:
        - active: Seleciona a aba ativa.
        - iter_rows(): Itera pelas linhas da planilha convertendo os valores em tuplas.
    """
    workbook = abrirPlanilhaExcel(caminhoArquivo)
    if workbook:
        planilha = workbook.active
        print("\n--- LISTA DE DESTINATÁRIOS CADASTRADOS ---")
        for linha in planilha.iter_rows(values_only=True):
            print(f"Nome: {linha[0]} | E-mail: {linha[1]} | Empresa: {linha[2]} | Mensagem: {linha[3]}")


def cadastrarDestinatario(caminhoArquivo, nome, email, empresa, mensagem):
    """
    Permite cadastrar um novo destinatário adicionando uma nova linha ao final da planilha.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo .xlsx.
        - nome (str): Nome do destinatário.
        - email (str): E-mail do destinatário.
        - empresa (str): Empresa do destinatário.
        - mensagem (str): Mensagem personalizada para o destinatário.
    """
    workbook = abrirPlanilhaExcel(caminhoArquivo)
    if workbook:
        planilha = workbook.active
        planilha.append([nome, email, empresa, mensagem])
        workbook.save(caminhoArquivo)
        print(f"{OK} Destinatário '{nome}' cadastrado com sucesso!")


def consultarDestinatarioPorEmail(caminhoArquivo, emailConsulta):
    """
    Permite consultar as informações de um destinatário buscando pelo seu endereço de e-mail.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo .xlsx.
        - emailConsulta (str): O e-mail que deseja localizar.
    """
    workbook = abrirPlanilhaExcel(caminhoArquivo)
    if workbook:
        planilha = workbook.active
        encontrado = False
        for linha in planilha.iter_rows(min_row=2, values_only=True):
            if linha[1].strip().lower() == emailConsulta.strip().lower():
                print(f"\n{OK} Destinatário Encontrado:")
                print(f"    - Nome: {linha[0]}")
                print(f"    - E-mail: {linha[1]}")
                print(f"    - Empresa: {linha[2]}")
                print(f"    - Mensagem: {linha[3]}")
                encontrado = True
                break
        if not encontrado:
            print(f"{ATENCAO} Nenhum destinatário foi encontrado com o e-mail: {emailConsulta}")


def alterarDestinatario(caminhoArquivo, emailAlvo):
    """
    Permite alterar informações de um destinatário existente com base no e-mail informado.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo .xlsx.
        - emailAlvo (str): E-mail do cadastro que sofrerá alteração.
    """
    workbook = abrirPlanilhaExcel(caminhoArquivo)
    if workbook:
        planilha = workbook.active
        encontrado = False
        
        for indice, linha in enumerate(planilha.iter_rows(min_row=2, values_only=True), start=2):
            if linha[1].strip().lower() == emailAlvo.strip().lower():
                print(f"{OK} Cadastro encontrado na linha {indice}. Insira os novos dados:")
                novoNome = input("Novo nome (ou Enter para manter): ")
                novoEmail = input("Novo e-mail (ou Enter para manter): ")
                novaEmpresa = input("Nova empresa (ou Enter para manter): ")
                novaMensagem = input("Nova mensagem (ou Enter para manter): ")
                
                if novoNome: planilha.cell(row=indice, column=1, value=novoNome)
                if novoEmail: planilha.cell(row=indice, column=2, value=novoEmail)
                if novaEmpresa: planilha.cell(row=indice, column=3, value=novaEmpresa)
                if novaMensagem: planilha.cell(row=indice, column=4, value=novaMensagem)
                
                workbook.save(caminhoArquivo)
                print(f"{OK} Informações atualizadas com sucesso!")
                encontrado = True
                break
        if not encontrado:
            print(f"{ATENCAO} E-mail '{emailAlvo}' não localizado para alteração.")


def removerDestinatario(caminhoArquivo, numeroLinha):
    """
    Permite remover um destinatário da base informando a linha correspondente.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo .xlsx.
        - numeroLinha (int): O número da linha na planilha a ser deletada.
    Métodos e Objetos usados:
        - delete_rows(): Remove uma linha específica da planilha pelo índice numérico.
    """
    workbook = abrirPlanilhaExcel(caminhoArquivo)
    if workbook:
        planilha = workbook.active
        if numeroLinha <= 1:
            print(f"{ERRO} Não é permitido remover o cabeçalho da planilha!")
            return
        try:
            planilha.delete_rows(numeroLinha)
            workbook.save(caminhoArquivo)
            print(f"{OK} Linha {numeroLinha} removida com sucesso da base.")
        except Exception as erro:
            print(f"{ERRO} Erro ao remover a linha: {erro}")


def enviarEmailIndividual(destinatarioEmail, assunto, corpo):
    """
    Envia um e-mail individual configurando uma conexão SMTP (ex: Gmail com App Password).
    Parâmetros:
        - destinatarioEmail (str): Endereço de e-mail de destino.
        - assunto (str): O assunto da mensagem.
        - corpo (str): O texto/corpo principal do e-mail.
    Métodos e Objetos usados:
        - MIMEMultipart(): Cria um container para mensagens compostas.
        - MIMEText(): Cria a parte de texto do e-mail.
        - smtplib.SMTP(): Conecta ao servidor SMTP na porta especificada (587).
        - starttls(): Inicia criptografia TLS.
        - login(): Realiza autenticação com credenciais.
        - sendmail(): Dispara o envio efetivo da mensagem.
    """
    remetenteEmail = input("Digite seu e-mail (remetente): ")
    remetenteSenha = input("Digite sua senha de aplicativo do e-mail: ")
    
    mensagem = MIMEMultipart()
    mensagem["From"] = remetenteEmail
    mensagem["To"] = destinatarioEmail
    mensagem["Subject"] = assunto
    
    mensagem.attach(MIMEText(corpo, "plain", "utf-8"))
    
    try:
        servidorSmtp = smtplib.SMTP("smtp.gmail.com", 587)
        servidorSmtp.starttls()
        servidorSmtp.login(remetenteEmail, remetenteSenha)
        servidorSmtp.sendmail(remetenteEmail, destinatarioEmail, mensagem.as_string())
        servidorSmtp.quit()
        print(f"{OK} E-mail enviado com sucesso para {destinatarioEmail}!")
        return True
    except Exception as erro:
        print(f"{ERRO} Erro ao enviar e-mail: {erro}")
        return False


def enviarEmailsEmMassa(caminhoArquivo):
    """
    Envia e-mails personalizados em lote para todos os destinatários cadastrados na planilha.
    Parâmetros:
        - caminhoArquivo (Path): Caminho do arquivo .xlsx contendo a base.
    """
    workbook = abrirPlanilhaExcel(caminhoArquivo)
    if not workbook:
        return
    
    planilha = workbook.active
    remetenteEmail = input("Digite seu e-mail (remetente): ")
    remetenteSenha = input("Digite sua senha de aplicativo do e-mail: ")
    
    try:
        servidorSmtp = smtplib.SMTP("smtp.gmail.com", 587)
        servidorSmtp.starttls()
        servidorSmtp.login(remetenteEmail, remetenteSenha)
        
        for linha in planilha.iter_rows(min_row=2, values_only=True):
            nome, emailDestino, empresa, textoPersonalizado = linha
            
            assunto = f"Mensagem Especial para {nome}"
            corpo = f"Olá {nome} da empresa {empresa}.\n\n{textoPersonalizado}\n\nAtenciosamente,\nSua Equipe"
            
            mensagem = MIMEMultipart()
            mensagem["From"] = remetenteEmail
            mensagem["To"] = emailDestino
            mensagem["Subject"] = assunto
            mensagem.attach(MIMEText(corpo, "plain", "utf-8"))
            
            servidorSmtp.sendmail(remetenteEmail, emailDestino, mensagem.as_string())
            print(f"{OK} E-mail enviado para: {emailDestino}")
            
        servidorSmtp.quit()
        print(f"{OK} Disparo em massa de e-mails concluído com sucesso!")
    except Exception as erro:
        print(f"{ERRO} Erro durante o envio em massa: {erro}")


def registrarHistorico(caminhoHistorico):
    """
    Registra o histórico dos envios realizados em um arquivo de texto de log.
    Parâmetros:
        - caminhoHistorico (Path): Caminho do arquivo de texto de histórico.
    Métodos e Objetos usados:
        - open(..., 'a'): Abre o arquivo no modo append (adicionar dados ao final).
        - datetime.now(): Obtém a data e hora atual do sistema.
    """
    try:
        destinatario = input("Digite o e-mail do destinatário: ")
        status = input("Digite o status do envio (Sucesso / Erro): ")
        erroEncontrado = input("Detalhes do erro (ou 'Nenhum'): ")
        
        dataHoraAtual = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        registro = f"[{dataHoraAtual}] - Destinatário: {destinatario} | Status: {status} | Erro: {erroEncontrado}\n"
        
        with open(caminhoHistorico, "a", encoding="utf-8") as arquivoLog:
            arquivoLog.write(registro)
            
        print(f"{OK} Histórico registrado com sucesso em '{caminhoHistorico.name}'.")
    except Exception as erro:
        print(f"{ERRO} Erro ao registrar o histórico: {erro}")