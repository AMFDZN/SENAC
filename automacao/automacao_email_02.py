
# pip install pyinstaller  

import os
import smtplib
import datetime
import getpass

from openpyxl import Workbook, load_workbook

from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase

from email import encoders




ARQUIVO_EXCEL = "destinatarios_email.xlsx"

CONFIG = "config.txt"

HISTORICO = "historico.txt"

SMTP_SERVER = "smtp.gmail.com"

SMTP_PORT = 587





def configurar_email():
    """
    Configura os dados do Gmail do usuário.

    Objetivo:
        Solicita ao usuário o endereço de Gmail
        e a senha de aplicativo.

    Processamento:
        Os dados informados são gravados
        no arquivo config.txt.

    Retorno:
        Não retorna valores.
    """
    

    print(

    )


    email = input(
        "Digite seu Gmail: "
    )


    senha = getpass.getpass(
        "Digite sua senha de aplicativo: "
    )


    arquivo = open(
        CONFIG,
        "w",
        encoding="utf-8"
    )


    arquivo.write(
        "email=" + email + "\n"
    )


    arquivo.write(
        "senha=" + senha + "\n"
    )


    arquivo.close()


    print(
        "\nConfiguração salva com sucesso!"
    )





def ler_config():

    if not os.path.exists(CONFIG):

        configurar_email()


    arquivo = open(
        CONFIG,
        "r",
        encoding="utf-8"
    )


    linhas = arquivo.readlines()


    email = linhas[0].replace(
        "email=",
        ""
    ).strip()


    senha = linhas[1].replace(
        "senha=",
        ""
    ).strip()


    arquivo.close()


    return email, senha







def criar_planilha():
    """
    Cria uma planilha Excel para cadastro
    dos destinatários.

    Objetivo:
        Criar o arquivo destinatarios_email.xlsx.

    Processamento:
        - Cria uma pasta de trabalho.
        - Cria uma aba chamada Clientes.
        - Adiciona os títulos das colunas.
        - Salva o arquivo.

    Retorno:
        Não retorna valores.
    """

    planilha = Workbook()


    aba = planilha.active


    aba.title = "Clientes"



    aba.append(
        [
            "Nome",
            "Email",
            "Empresa",
            "Mensagem",
            "Anexo"
        ]
    )


    aba.append(
        [
            "AMFDZN",
            "amfdzn@gmail.com",
            "Ace Lio Design",
            "Olá, temos uma proposta\nteste de quebra de linha.",
            ""
        ]
    )


    aba.append(
        [
            "Amf Cadastros",
            "amf.cadastros@gmail.com",
            "Recebedora de Spams Ltda",
            "Novo catálogo disponível.",
            ""
        ]
    )


    planilha.save(
        ARQUIVO_EXCEL
    )


    print(
        "Planilha criada!"
    )





def abrir_planilha():


    if not os.path.exists(
        ARQUIVO_EXCEL
    ):


        print(
            "Planilha não encontrada!"
        )


        return None



    return load_workbook(
        ARQUIVO_EXCEL
    )





def listar_destinatarios():


    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha["Clientes"]


    print(
        "\nLISTA DE CLIENTES\n"
    )


    for linha in aba.iter_rows(
        values_only=True
    ):

        print(
            linha
        )








def registrar_historico(
        email,
        status
):


    arquivo = open(
        HISTORICO,
        "a",
        encoding="utf-8"
    )


    data = datetime.datetime.now()


    arquivo.write(
        f"{data} - {email} - {status}\n"
    )


    arquivo.close()








def validar_email(email):


    if "@" in email:

        return True


    return False





def enviar_email(nome, email, empresa, texto, anexo=None):
    """
    Envia um email personalizado.

    Parâmetros:
        nome:
            Nome do destinatário.

        email:
            Endereço de email.

        empresa:
            Empresa do contato.

        texto:
            Mensagem que será enviada.

        anexo:
            Caminho do arquivo anexado.

    Retorno:
        Não retorna valores.
    """

    usuario, senha = ler_config()

    if not validar_email(email):
        registrar_historico(
            email,
            "EMAIL INVALIDO"
        )

        print("Email inválido:", email)
        return


    mensagem = MIMEMultipart()

    mensagem["From"] = usuario
    mensagem["To"] = email
    mensagem["Subject"] = "Contato comercial"


    corpo = f"""
Olá {nome},

Empresa:
{empresa}

Mensagem:
{texto}


Obrigado.
"""


    mensagem.attach(
        MIMEText(
            corpo,
            "plain",
            "utf-8"
        )
    )


    if anexo and os.path.exists(anexo):

        with open(anexo, "rb") as arquivo:

            parte = MIMEBase(
                "application",
                "octet-stream"
            )

            parte.set_payload(
                arquivo.read()
            )


        encoders.encode_base64(parte)


        parte.add_header(
            "Content-Disposition",
            f'attachment; filename="{os.path.basename(anexo)}"'
        )


        mensagem.attach(parte)


    try:

        with smtplib.SMTP(
            SMTP_SERVER,
            SMTP_PORT
        ) as servidor:

            servidor.starttls()

            servidor.login(
                usuario,
                senha
            )

            servidor.sendmail(
                usuario,
                email,
                mensagem.as_string()
            )


        registrar_historico(
            email,
            "ENVIADO"
        )

        print(
            "Enviado:",
            email
        )


    except Exception as erro:

        registrar_historico(
            email,
            "ERRO"
        )

        print(
            "Erro:",
            erro
        )




def enviar_todos():

    planilha = abrir_planilha()

    if planilha is None:
        return


    aba = planilha["Clientes"]

    contador = 0


    for linha in aba.iter_rows(
        min_row=2,
        values_only=True
    ):

       
        if not linha[0]:
            continue


        nome = linha[0]
        email = linha[1]
        empresa = linha[2]
        texto = linha[3]
        anexo = linha[4]


        enviar_email(
            nome,
            email,
            empresa,
            texto,
            anexo
        )


        contador += 1


    print(
        "\nTotal processado:",
        contador
    )







def gerar_executavel():


    print(
        "\nGerando executável...\n"
    )



    comando = (

        "pyinstaller "

        "--onefile "

        "--name SistemaEmail "

        "automacao_email_02.py"

    )



    os.system(
        comando
    )



    pasta = os.path.abspath(
        "dist"
    )



    print(
        """
Executável criado!

Local:
"""
    )


    print(
        pasta
    )


    os.startfile(
        pasta
    )








def menu():


    while True:


        print(
        """

1 - Configurar Gmail

2 - Criar planilha

3 - Listar destinatários

4 - Enviar emails

5 - Gerar executável

0 - Sair


"""
        )



        opcao = input(
            "Escolha uma opção: "
        )



        match opcao:


            case "1":

                configurar_email()



            case "2":

                criar_planilha()



            case "3":

                listar_destinatarios()



            case "4":

                enviar_todos()



            case "5":

                gerar_executavel()



            case "0":

                print(
                    "Programa encerrado."
                )

                break



            case _:

                print(
                    "Opção inválida!"
                )








if __name__ == "__main__":

    menu()