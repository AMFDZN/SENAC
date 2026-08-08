#pip install openpyxl

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment
import os


ARQUIVO = "planilha_openpyxl.xlsx"



def criar_planilha():
    """
    Cria uma nova planilha Excel utilizando openpyxl.

    Utiliza:

        Workbook() -> Cria um novo arquivo Excel.

        save() -> Salva o arquivo no disco.

    Conceitos:

        Workbook:
            Representa o arquivo Excel.

        Worksheet:
            Representa uma aba dentro do arquivo.

    """

    planilha = Workbook()

    aba = planilha.active

    aba.title = "Produtos"


    aba["A1"] = "Produto"
    aba["B1"] = "Quantidade"
    aba["C1"] = "Preço"


    aba.append(
        [
            "Notebook",
            10,
            3500
        ]
    )


    aba.append(
        [
            "Mouse",
            25,
            80
        ]
    )


    aba.append(
        [
            "Teclado",
            15,
            150
        ]
    )


    planilha.save(ARQUIVO)


    print("\nPlanilha criada com sucesso!")





def verificar_arquivo():
    """
    Verifica se a planilha existe.

    Retorno:

        True:
            Arquivo encontrado.

        False:
            Arquivo inexistente.

    """

    if not os.path.exists(ARQUIVO):

        print("\nA planilha não existe.")

        print("Crie a planilha primeiro (Opção 1).")

        return False


    return True





def abrir_planilha():
    """
    Abre uma planilha existente.

    Utiliza:

        load_workbook()
            Abre arquivos Excel existentes.

    Retorno:

        Workbook carregado.

    """

    if not verificar_arquivo():

        return None


    planilha = load_workbook(ARQUIVO)


    return planilha





def listar_planilha():
    """
    Mostra todos os dados da planilha.

    Utiliza:

        iter_rows()
            Percorre todas as linhas da planilha.

        values_only=True
            Retorna somente os valores das células.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return


    aba = planilha.active


    print("\nPLANILHA\n")


    for linha in aba.iter_rows(values_only=True):

        print(linha)





def escrever_celula():
    """
    Escreve um valor em uma célula específica.

    Utiliza:

        Cell
            Representa uma célula da planilha.

        value
            Define o conteúdo da célula.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return


    aba = planilha.active


    celula = input("Informe a célula (ex: A1): ")


    valor = input("Digite o valor: ")



    aba[celula] = valor



    planilha.save(ARQUIVO)



    print("Valor inserido com sucesso.")





def ler_celula():
    """
    Realiza a leitura de uma célula.

    Utiliza:

        worksheet[célula]
            Acessa uma célula específica.

        value
            Retorna o conteúdo armazenado.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return


    aba = planilha.active



    celula = input("Informe a célula (ex: A1): ")



    valor = aba[celula].value



    print("\nValor encontrado:")

    print(valor)





def alterar_celula():
    """
    Altera o conteúdo de uma célula existente.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return


    aba = planilha.active



    celula = input("Informe a célula: ")


    novo_valor = input("Novo valor: ")



    aba[celula] = novo_valor



    planilha.save(ARQUIVO)



    print("Célula alterada com sucesso.")





def adicionar_linha():
    """
    Adiciona uma nova linha na planilha.

    Utiliza:

        append()
            Insere dados na próxima linha disponível.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return


    aba = planilha.active



    produto = input("Produto: ")

    quantidade = int(input("Quantidade: "))

    preco = float(input("Preço: "))



    aba.append(
        [
            produto,
            quantidade,
            preco
        ]
    )



    planilha.save(ARQUIVO)



    print("Linha adicionada com sucesso.")



def remover_linha():
    """
    Remove uma linha da planilha.

    Utiliza:

        delete_rows()
            Remove uma ou mais linhas.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return


    aba = planilha.active


    linha = int(
        input("Informe o número da linha que deseja remover: ")
    )


    if linha < 1 or linha > aba.max_row:

        print("Linha inexistente.")

        return


    aba.delete_rows(linha)


    planilha.save(ARQUIVO)


    print("Linha removida com sucesso.")





def adicionar_coluna():
    """
    Adiciona uma nova coluna na planilha.

    Utiliza:

        insert_cols()
            Insere colunas vazias.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return


    aba = planilha.active


    coluna = int(
        input("Informe a posição da coluna: ")
    )


    aba.insert_cols(coluna)


    titulo = input(
        "Digite o título da nova coluna: "
    )


    aba.cell(
        row=1,
        column=coluna
    ).value = titulo



    planilha.save(ARQUIVO)


    print("Coluna adicionada com sucesso.")





def remover_coluna():
    """
    Remove uma coluna da planilha.

    Utiliza:

        delete_cols()
            Remove colunas existentes.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return


    aba = planilha.active


    coluna = int(
        input("Informe o número da coluna: ")
    )


    if coluna < 1 or coluna > aba.max_column:

        print("Coluna inexistente.")

        return



    aba.delete_cols(coluna)


    planilha.save(ARQUIVO)


    print("Coluna removida com sucesso.")





def criar_aba():
    """
    Cria uma nova aba dentro do arquivo Excel.

    Utiliza:

        create_sheet()
            Cria uma nova Worksheet.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    nome = input(
        "Nome da nova aba: "
    )


    if nome in planilha.sheetnames:

        print("Essa aba já existe.")

        return



    planilha.create_sheet(nome)



    planilha.save(ARQUIVO)



    print("Aba criada com sucesso.")





def listar_abas():
    """
    Lista todas as abas existentes.

    Utiliza:

        sheetnames
            Retorna os nomes das Worksheets.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    print("\nABAS EXISTENTES\n")


    for aba in planilha.sheetnames:

        print("-", aba)





def selecionar_aba():
    """
    Permite selecionar uma aba específica.

    Retorna:

        Worksheet selecionada.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return None



    print("\nAbas disponíveis:")


    for aba in planilha.sheetnames:

        print("-", aba)



    nome = input(
        "Informe o nome da aba: "
    )



    if nome not in planilha.sheetnames:

        print("Aba não encontrada.")

        return None



    return planilha[nome]





def renomear_aba():
    """
    Altera o nome de uma aba.

    Utiliza:

        title
            Propriedade que representa o nome da Worksheet.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    antiga = input(
        "Nome da aba atual: "
    )


    if antiga not in planilha.sheetnames:

        print("Aba inexistente.")

        return



    nova = input(
        "Novo nome: "
    )



    planilha[antiga].title = nova



    planilha.save(ARQUIVO)



    print("Aba renomeada com sucesso.")





def remover_aba():
    """
    Remove uma aba do arquivo Excel.

    Utiliza:

        remove()
            Exclui uma Worksheet.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    nome = input(
        "Nome da aba que deseja remover: "
    )



    if nome not in planilha.sheetnames:

        print("Aba inexistente.")

        return



    aba = planilha[nome]


    planilha.remove(aba)



    planilha.save(ARQUIVO)



    print("Aba removida com sucesso.")





def mostrar_dimensoes():
    """
    Mostra o tamanho utilizado da planilha.

    Utiliza:

        max_row
            Quantidade de linhas utilizadas.

        max_column
            Quantidade de colunas utilizadas.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    print("\nDimensões da planilha\n")


    print(
        "Quantidade de linhas:",
        aba.max_row
    )


    print(
        "Quantidade de colunas:",
        aba.max_column
    )





def criar_formula():
    """
    Insere uma fórmula Excel em uma célula.

    Exemplo:

        Total = Quantidade * Preço


    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    linha = int(
        input("Informe a linha: ")
    )



    aba[f"D{linha}"] = (
        f"=B{linha}*C{linha}"
    )



    planilha.save(ARQUIVO)



    print("Fórmula adicionada.")





def alterar_aba_atual():
    """
    Permite trocar a aba utilizada
    para as operações.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    nome = input(
        "Informe a aba: "
    )


    if nome in planilha.sheetnames:

        planilha.active = (
            planilha.sheetnames.index(nome)
        )

        planilha.save(ARQUIVO)

        print("Aba alterada.")

    else:

        print("Aba não encontrada.")










def aplicar_negrito():
    """
    Aplica negrito em uma célula.

    Utiliza:

        Font()
            Define características da fonte.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active


    celula = input(
        "Informe a célula: "
    )


    aba[celula].font = Font(
        bold=True
    )


    planilha.save(ARQUIVO)


    print("Negrito aplicado.")





def alterar_fonte():
    """
    Altera tamanho e cor da fonte.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    celula = input(
        "Informe a célula: "
    )


    tamanho = int(
        input("Tamanho da fonte: ")
    )


    cor = input(
        "Cor hexadecimal (ex: FF0000): "
    )



    aba[celula].font = Font(
        size=tamanho,
        color=cor
    )


    planilha.save(ARQUIVO)



    print("Fonte alterada.")





def aplicar_cor_fundo():
    """
    Aplica cor de preenchimento em uma célula.

    Utiliza:

        PatternFill()
            Define preenchimento.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    celula = input(
        "Informe a célula: "
    )


    cor = input(
        "Cor hexadecimal: "
    )


    aba[celula].fill = PatternFill(
        "solid",
        fgColor=cor
    )



    planilha.save(ARQUIVO)



    print("Cor aplicada.")





def aplicar_borda():
    """
    Adiciona borda em uma célula.

    Utiliza:

        Border()
        Side()

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    celula = input(
        "Informe a célula: "
    )


    lado = Side(
        style="thin",
        color="000000"
    )


    aba[celula].border = Border(
        left=lado,
        right=lado,
        top=lado,
        bottom=lado
    )


    planilha.save(ARQUIVO)



    print("Borda aplicada.")





def centralizar_celula():
    """
    Centraliza o conteúdo de uma célula.

    Utiliza:

        Alignment()

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    celula = input(
        "Informe a célula: "
    )


    aba[celula].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


    planilha.save(ARQUIVO)



    print("Alinhamento aplicado.")





def ajustar_colunas():
    """
    Ajusta automaticamente a largura
    das colunas conforme o conteúdo.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    for coluna in aba.columns:


        maior = 0


        letra = coluna[0].column_letter



        for celula in coluna:


            if celula.value:


                tamanho = len(
                    str(celula.value)
                )


                if tamanho > maior:

                    maior = tamanho



        aba.column_dimensions[letra].width = maior + 2



    planilha.save(ARQUIVO)



    print("Colunas ajustadas.")





def mesclar_celulas():
    """
    Mescla células.

    Exemplo:

        A1:D1

    Muito utilizado para títulos.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    intervalo = input(
        "Informe intervalo (A1:D1): "
    )



    aba.merge_cells(
        intervalo
    )



    planilha.save(ARQUIVO)



    print("Células mescladas.")





def congelar_cabecalho():
    """
    Congela linhas superiores.

    Usado em relatórios grandes.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    aba.freeze_panes = "A2"



    planilha.save(ARQUIVO)



    print("Cabeçalho congelado.")





def aplicar_filtro():
    """
    Cria filtro automático na tabela.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    aba.auto_filter.ref = aba.dimensions



    planilha.save(ARQUIVO)



    print("Filtro aplicado.")





def criar_tabela_excel():
    """
    Cria uma tabela Excel formatada.

    Utiliza:

        Table()

        TableStyleInfo()

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    tabela = Table(
        displayName="TabelaProdutos",
        ref=aba.dimensions
    )


    estilo = TableStyleInfo(
        name="TableStyleMedium9",
        showRowStripes=True,
        showColumnStripes=False
    )


    tabela.tableStyleInfo = estilo



    aba.add_table(
        tabela
    )



    planilha.save(ARQUIVO)



    print("Tabela criada.")





def criar_grafico():
    """
    Cria gráfico de barras.

    Utiliza:

        BarChart()

        Reference()

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    grafico = BarChart()



    dados = Reference(
        aba,
        min_col=2,
        min_row=1,
        max_row=aba.max_row
    )


    categorias = Reference(
        aba,
        min_col=1,
        min_row=2,
        max_row=aba.max_row
    )



    grafico.add_data(
        dados,
        titles_from_data=True
    )


    grafico.set_categories(
        categorias
    )



    grafico.title = "Relatório"



    aba.add_chart(
        grafico,
        "F2"
    )



    planilha.save(ARQUIVO)



    print("Gráfico criado.")





def adicionar_comentario():
    """
    Adiciona comentário em uma célula.

    """

    planilha = abrir_planilha()


    if planilha is None:

        return



    aba = planilha.active



    celula = input(
        "Informe a célula: "
    )


    texto = input(
        "Comentário: "
    )



    aba[celula].comment = Comment(
        texto,
        "Python"
    )


    planilha.save(ARQUIVO)



    print("Comentário adicionado.")





def menu():

    while True:


        print("""
        
1  - Criar planilha
2  - Listar planilha
3  - Escrever célula
4  - Ler célula
5  - Alterar célula
6  - Adicionar linha
7  - Remover linha

8  - Adicionar coluna
9  - Remover coluna

10 - Criar aba
11 - Listar abas
12 - Renomear aba
13 - Remover aba

14 - Mostrar dimensões
15 - Criar fórmula

16 - Aplicar negrito
17 - Alterar fonte
18 - Cor de fundo
19 - Aplicar borda
20 - Centralizar célula

21 - Ajustar colunas
22 - Mesclar células
23 - Congelar cabeçalho
24 - Aplicar filtro
25 - Criar tabela Excel
26 - Criar gráfico
27 - Adicionar comentário


0 - Sair

""")


        opcao = input(
            "Escolha uma opção: "
        )


        match opcao:


            case "1":

                criar_planilha()
                input("Pressione Enter para voltar ao menu.")


            case "2":

                listar_planilha()
                input("Pressione Enter para voltar ao menu.")


            case "3":

                escrever_celula()
                input("Pressione Enter para voltar ao menu.")


            case "4":

                ler_celula()
                input("Pressione Enter para voltar ao menu.")


            case "5":

                alterar_celula()
                input("Pressione Enter para voltar ao menu.")


            case "6":

                adicionar_linha()
                input("Pressione Enter para voltar ao menu.")


            case "7":

                remover_linha()
                input("Pressione Enter para voltar ao menu.")


            case "8":

                adicionar_coluna()
                input("Pressione Enter para voltar ao menu.")


            case "9":

                remover_coluna()
                input("Pressione Enter para voltar ao menu.")


            case "10":

                criar_aba()
                input("Pressione Enter para voltar ao menu.")


            case "11":

                listar_abas()
                input("Pressione Enter para voltar ao menu.")


            case "12":

                renomear_aba()
                input("Pressione Enter para voltar ao menu.")


            case "13":

                remover_aba()
                input("Pressione Enter para voltar ao menu.")


            case "14":

                mostrar_dimensoes()
                input("Pressione Enter para voltar ao menu.")


            case "15":

                criar_formula()
                input("Pressione Enter para voltar ao menu.")


            case "16":

                aplicar_negrito()
                input("Pressione Enter para voltar ao menu.")


            case "17":

                alterar_fonte()
                input("Pressione Enter para voltar ao menu.")


            case "18":

                aplicar_cor_fundo()
                input("Pressione Enter para voltar ao menu.")


            case "19":

                aplicar_borda()
                input("Pressione Enter para voltar ao menu.")


            case "20":

                centralizar_celula()
                input("Pressione Enter para voltar ao menu.")


            case "21":

                ajustar_colunas()
                input("Pressione Enter para voltar ao menu.")


            case "22":

                mesclar_celulas()
                input("Pressione Enter para voltar ao menu.")


            case "23":

                congelar_cabecalho()
                input("Pressione Enter para voltar ao menu.")


            case "24":

                aplicar_filtro()
                input("Pressione Enter para voltar ao menu.")


            case "25":

                criar_tabela_excel()
                input("Pressione Enter para voltar ao menu.")


            case "26":

                criar_grafico()
                input("Pressione Enter para voltar ao menu.")


            case "27":

                adicionar_comentario()
                input("Pressione Enter para voltar ao menu.")


            case "0":

                print(
                    "Programa encerrado."
                )

                break


            case _:

                print("Opção inválida.")
                input("Pressione Enter para voltar ao menu.")





if __name__ == "__main__":

    menu()